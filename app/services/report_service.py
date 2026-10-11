from pathlib import Path
from typing import Dict, List
from datetime import datetime
import pandas as pd
from reportlab.lib.pagesizes import letter, A4
from reportlab.lib import colors
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer, PageBreak
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import inch
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side

from app.storage.json_store import invoices_store, vendors_store, anomalies_store
from app.core.config import settings
from app.services.duplicate_service import duplicate_service

class ReportService:
    """Service for generating PDF and Excel reports"""
    
    def generate_monthly_report(self, month: str = None) -> Dict[str, str]:
        """Generate monthly report in PDF and Excel formats"""
        if not month:
            month = datetime.now().strftime("%Y-%m")
        
        # Get data
        invoices = invoices_store.read("invoices.json", {})
        vendors = vendors_store.read("vendors.json", {})
        anomalies = anomalies_store.read("anomalies.json", {})
        duplicates = duplicate_service.find_all_duplicates()
        
        # Filter invoices for the month
        month_invoices = [
            inv for inv in invoices.values()
            if inv.get("invoice_date", "").startswith(month)
        ]
        
        # Generate report data
        report_data = {
            "month": month,
            "total_invoices": len(month_invoices),
            "total_amount": sum(inv.get("total_amount", 0) or 0 for inv in month_invoices),
            "invoices": month_invoices,
            "vendors": vendors,
            "anomalies": anomalies,
            "duplicates": duplicates
        }
        
        # Generate PDF report
        pdf_path = self._generate_pdf_report(report_data)
        
        # Generate Excel report
        excel_path = self._generate_excel_report(report_data)
        
        return {
            "month": month,
            "pdf_path": str(pdf_path),
            "excel_path": str(excel_path)
        }
    
    def _generate_pdf_report(self, data: Dict) -> Path:
        """Generate PDF report"""
        filename = f"monthly_report_{data['month']}.pdf"
        filepath = settings.REPORTS_DIR / filename
        
        doc = SimpleDocTemplate(str(filepath), pagesize=A4)
        styles = getSampleStyleSheet()
        story = []
        
        # Title
        title_style = ParagraphStyle(
            'CustomTitle',
            parent=styles['Heading1'],
            fontSize=24,
            textColor=colors.darkblue,
            spaceAfter=30
        )
        story.append(Paragraph(f"Monthly Invoice Report - {data['month']}", title_style))
        story.append(Spacer(1, 12))
        
        # Summary
        summary_style = ParagraphStyle(
            'CustomSummary',
            parent=styles['Normal'],
            fontSize=12,
            spaceAfter=20
        )
        
        story.append(Paragraph(f"<b>Total Invoices:</b> {data['total_invoices']}", summary_style))
        story.append(Paragraph(f"<b>Total Amount:</b> ₹{data['total_amount']:,.2f}", summary_style))
        story.append(Spacer(1, 20))
        
        # Invoice Table
        if data['invoices']:
            story.append(Paragraph("<b>Invoices</b>", styles['Heading2']))
            story.append(Spacer(1, 12))
            
            table_data = [['Invoice ID', 'Vendor', 'Date', 'Amount', 'Status']]
            for inv in data['invoices'][:20]:  # Limit to 20 invoices
                table_data.append([
                    inv.get('id', ''),
                    inv.get('vendor_name', 'Unknown'),
                    inv.get('invoice_date', ''),
                    f"₹{inv.get('total_amount', 0):,.2f}",
                    inv.get('status', '')
                ])
            
            table = Table(table_data, colWidths=[1.5*inch, 2*inch, 1*inch, 1*inch, 1*inch])
            table.setStyle(TableStyle([
                ('BACKGROUND', (0, 0), (-1, 0), colors.grey),
                ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
                ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
                ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
                ('FONTSIZE', (0, 0), (-1, 0), 10),
                ('BOTTOMPADDING', (0, 0), (-1, 0), 12),
                ('BACKGROUND', (0, 1), (-1, -1), colors.beige),
                ('GRID', (0, 0), (-1, -1), 1, colors.black)
            ]))
            story.append(table)
            story.append(Spacer(1, 20))
        
        # Vendor Summary
        if data['vendors']:
            story.append(Paragraph("<b>Vendor Summary</b>", styles['Heading2']))
            story.append(Spacer(1, 12))
            
            vendor_data = [['Vendor', 'Total Invoices', 'Total Spending']]
            for vendor in list(data['vendors'].values())[:10]:
                vendor_data.append([
                    vendor.get('name', 'Unknown'),
                    str(vendor.get('total_invoices', 0)),
                    f"₹{vendor.get('total_spending', 0):,.2f}"
                ])
            
            table = Table(vendor_data, colWidths=[2*inch, 1.5*inch, 1.5*inch])
            table.setStyle(TableStyle([
                ('BACKGROUND', (0, 0), (-1, 0), colors.grey),
                ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
                ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
                ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
                ('FONTSIZE', (0, 0), (-1, 0), 10),
                ('BOTTOMPADDING', (0, 0), (-1, 0), 12),
                ('BACKGROUND', (0, 1), (-1, -1), colors.beige),
                ('GRID', (0, 0), (-1, -1), 1, colors.black)
            ]))
            story.append(table)
            story.append(Spacer(1, 20))
        
        # Category Breakdown
        category_totals = {}
        for inv in data['invoices']:
            cat = inv.get('category', 'Other')
            category_totals[cat] = category_totals.get(cat, 0) + (inv.get('total_amount', 0) or 0)
        
        if category_totals:
            story.append(Paragraph("<b>Category Breakdown</b>", styles['Heading2']))
            story.append(Spacer(1, 10))
            cat_table_data = [['Category', 'Total Amount']]
            for c_name, c_amt in category_totals.items():
                cat_table_data.append([c_name, f"₹{c_amt:,.2f}"])
            cat_table = Table(cat_table_data, colWidths=[3*inch, 2*inch])
            cat_table.setStyle(TableStyle([
                ('BACKGROUND', (0, 0), (-1, 0), colors.navy),
                ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
                ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
                ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
                ('GRID', (0, 0), (-1, -1), 1, colors.black)
            ]))
            story.append(cat_table)
            story.append(Spacer(1, 20))

        # Anomalies
        if data['anomalies']:
            story.append(Paragraph("<b>Anomalies Detected</b>", styles['Heading2']))
            story.append(Spacer(1, 12))
            
            anomaly_count = len(data['anomalies'])
            story.append(Paragraph(f"Total anomalies detected: {anomaly_count}", summary_style))
            story.append(Spacer(1, 15))

        # Duplicates
        if data.get('duplicates'):
            story.append(Paragraph("<b>Potential Duplicates</b>", styles['Heading2']))
            story.append(Spacer(1, 10))
            dup_count = len(data['duplicates'])
            story.append(Paragraph(f"Potential duplicate groups detected: {dup_count}", summary_style))
            story.append(Spacer(1, 15))
        
        # Build PDF
        doc.build(story)
        
        return filepath
    
    def _generate_excel_report(self, data: Dict) -> Path:
        """Generate Excel report with multiple sheets"""
        filename = f"monthly_report_{data['month']}.xlsx"
        filepath = settings.REPORTS_DIR / filename
        
        wb = Workbook()
        
        # Remove default sheet
        wb.remove(wb.active)
        
        # Summary Sheet
        ws_summary = wb.create_sheet("Summary")
        ws_summary['A1'] = f"Monthly Report - {data['month']}"
        ws_summary['A1'].font = Font(bold=True, size=14)
        ws_summary['A3'] = "Total Invoices"
        ws_summary['B3'] = data['total_invoices']
        ws_summary['A4'] = "Total Amount"
        ws_summary['B4'] = data['total_amount']
        
        # Invoices Sheet
        if data['invoices']:
            ws_invoices = wb.create_sheet("Invoices")
            headers = ['Invoice ID', 'Vendor', 'Invoice Number', 'Date', 'Amount', 'Category', 'Status']
            ws_invoices.append(headers)
            
            # Style headers
            for cell in ws_invoices[1]:
                cell.font = Font(bold=True)
                cell.fill = PatternFill(start_color="CCCCCC", end_color="CCCCCC", fill_type="solid")
            
            for inv in data['invoices']:
                row = [
                    inv.get('id', ''),
                    inv.get('vendor_name', ''),
                    inv.get('invoice_number', ''),
                    inv.get('invoice_date', ''),
                    inv.get('total_amount', 0),
                    inv.get('category', ''),
                    inv.get('status', '')
                ]
                ws_invoices.append(row)
        
        # Vendors Sheet
        if data['vendors']:
            ws_vendors = wb.create_sheet("Vendors")
            headers = ['Vendor', 'Total Invoices', 'Total Spending', 'Average Invoice', 'Categories']
            ws_vendors.append(headers)
            
            for cell in ws_vendors[1]:
                cell.font = Font(bold=True)
                cell.fill = PatternFill(start_color="CCCCCC", end_color="CCCCCC", fill_type="solid")
            
            for vendor in data['vendors'].values():
                row = [
                    vendor.get('name', ''),
                    vendor.get('total_invoices', 0),
                    vendor.get('total_spending', 0),
                    vendor.get('average_invoice', 0),
                    ', '.join(vendor.get('categories', []))
                ]
                ws_vendors.append(row)
        
        # Categories Sheet
        ws_categories = wb.create_sheet("Categories")
        ws_categories['A1'] = "Category Breakdown"
        ws_categories['A1'].font = Font(bold=True, size=12)
        
        # Calculate category totals
        category_totals = {}
        for inv in data['invoices']:
            cat = inv.get('category', 'Other')
            amount = inv.get('total_amount', 0) or 0
            if cat not in category_totals:
                category_totals[cat] = 0
            category_totals[cat] += amount
        
        ws_categories['A3'] = "Category"
        ws_categories['B3'] = "Total Amount"
        for cell in ws_categories[3]:
            cell.font = Font(bold=True)
        
        row = 4
        for cat, amount in category_totals.items():
            ws_categories[f'A{row}'] = cat
            ws_categories[f'B{row}'] = amount
            row += 1
        
        # Anomalies Sheet
        if data['anomalies']:
            ws_anomalies = wb.create_sheet("Anomalies")
            headers = ['Invoice ID', 'Anomalies', 'Detected At']
            ws_anomalies.append(headers)
            
            for cell in ws_anomalies[1]:
                cell.font = Font(bold=True)
                cell.fill = PatternFill(start_color="CCCCCC", end_color="CCCCCC", fill_type="solid")
            
            for anomaly_id, anomaly_data in data['anomalies'].items():
                row = [
                    anomaly_id,
                    ', '.join(anomaly_data.get('anomalies', [])),
                    anomaly_data.get('detected_at', '')
                ]
                ws_anomalies.append(row)
        
        # Duplicates Sheet
        if data.get('duplicates'):
            ws_duplicates = wb.create_sheet("Duplicates")
            headers = ['Duplicate Type', 'Invoice 1', 'Invoice 2', 'Details']
            ws_duplicates.append(headers)
            for cell in ws_duplicates[1]:
                cell.font = Font(bold=True)
                cell.fill = PatternFill(start_color="CCCCCC", end_color="CCCCCC", fill_type="solid")
            for dup in data['duplicates']:
                ws_duplicates.append([
                    dup.get('type', 'POSSIBLE_DUPLICATE'),
                    dup.get('invoice_id_1', ''),
                    dup.get('invoice_id_2', ''),
                    dup.get('reason', '')
                ])

        # Save workbook
        wb.save(str(filepath))
        
        return filepath
    
    def generate_csv_export(self) -> Path:
        """Generate CSV export of all invoices"""
        invoices = invoices_store.read("invoices.json", {})
        
        if not invoices:
            raise Exception("No invoices to export")
        
        # Convert to DataFrame
        df = pd.DataFrame(list(invoices.values()))
        
        # Select relevant columns
        columns_to_export = [
            'id', 'vendor_name', 'invoice_number', 'invoice_date', 'due_date',
            'subtotal', 'tax_amount', 'total_amount', 'category', 'status',
            'confidence_score', 'duplicate_status'
        ]
        
        df_export = df[columns_to_export]
        
        # Save to CSV
        filename = f"invoices_export_{datetime.now().strftime('%Y%m%d_%H%M%S')}.csv"
        filepath = settings.REPORTS_DIR / filename
        df_export.to_csv(filepath, index=False)
        
        return filepath

report_service = ReportService()
